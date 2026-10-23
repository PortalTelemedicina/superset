# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.
"""Leitura de CSV e XLSX da importação de usuários. Tudo vira texto."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from typing import Any, Iterable, Optional

from superset.extensions.portal.services.user_import.models import (
    MAX_CELL_LEN,
    ParsedFile,
    RawRow,
    WARNING_USERNAME_NUMERICO_XLSX,
)

DEFAULT_MAX_ROWS = 500
DEFAULT_MAX_BYTES = 1024 * 1024
# soma dos tamanhos descomprimidos do XLSX; um arquivo pequeno pode ser zip bomb
DEFAULT_XLSX_MAX_UNCOMPRESSED = 20 * 1024 * 1024
XLSX_MAX_ENTRIES = 200
XLSX_MAX_COLUMNS = 50
# Limite físico do Excel. Célula só formatada lá embaixo faz o openpyxl devolver
# linhas vazias até ela; o teto de linhas de dados vale só para linhas não vazias.
XLSX_MAX_SCAN_LINES = 1_048_576

# código de erro -> a camada de API decide o HTTP (400 ou 422)
ERRO_EXTENSAO = "extensao_invalida"
ERRO_COLUNAS = "colunas_ausentes"
ERRO_VAZIO = "arquivo_vazio"
ERRO_ENCODING = "encoding_invalido"
ERRO_LIMITE_LINHAS = "limite_linhas"
ERRO_LIMITE_BYTES = "limite_bytes"
ERRO_ILEGIVEL = "arquivo_ilegivel"

_SINONIMOS = {
    "nome": "nome",
    "usuario": "usuario",
    "cpf": "usuario",
    "username": "usuario",
    "senha": "senha",
    "email": "email",
    "role": "role",
    "perfil": "role",
}
_OBRIGATORIAS = (
    ("nome", "NOME"),
    ("usuario", "USUARIO"),
    ("senha", "SENHA"),
    ("email", "EMAIL"),
)


class ParseError(Exception):
    """Erro de leitura. A mensagem nunca ecoa conteúdo do arquivo."""

    def __init__(
        self, code: str, message: str, missing_columns: tuple[str, ...] = ()
    ) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.missing_columns = missing_columns


def normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", "" if value is None else str(value))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", text.lower())


def _map_header(cells: Iterable[Any]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, cell in enumerate(cells):
        field = _SINONIMOS.get(normalize_header(cell))
        if field and field not in mapping:
            mapping[field] = index
    missing = tuple(label for key, label in _OBRIGATORIAS if key not in mapping)
    if missing:
        raise ParseError(
            ERRO_COLUNAS, "Colunas obrigatórias ausentes", missing_columns=missing
        )
    return mapping


def _cell_text(cells: list[Any], mapping: dict[str, int], field: str) -> str:
    index = mapping.get(field)
    if index is None or index >= len(cells):
        return ""
    value = cells[index]
    return "" if value is None else str(value)


def _build_row(
    line: int, cells: list[Any], mapping: dict[str, int], username: Optional[str] = None
) -> RawRow:
    values = {f: _cell_text(cells, mapping, f) for f in mapping}
    if username is not None:
        values["usuario"] = username
    longa = any(len(v) > MAX_CELL_LEN for v in values.values())
    if longa:
        values = {k: v[:MAX_CELL_LEN] for k, v in values.items()}
    return RawRow(
        line=line,
        nome=values.get("nome", ""),
        usuario=values.get("usuario", ""),
        senha=values.get("senha", ""),
        email=values.get("email", ""),
        role=values.get("role", ""),
        celula_longa=longa,
    )


def _decode_csv(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    try:
        return content.decode("cp1252")
    except UnicodeDecodeError as ex:
        raise ParseError(ERRO_ENCODING, "Codificação do arquivo não reconhecida") from ex


def _detect_delimiter(text: str) -> str:
    first = next((ln for ln in text.splitlines() if ln.strip()), "")
    best, best_count = ",", -1
    for delimiter in (",", ";"):
        try:
            count = len(next(csv.reader([first], delimiter=delimiter), []))
        except csv.Error:
            count = 0
        if count > best_count:
            best, best_count = delimiter, count
    return best


def parse_csv(content: bytes, max_rows: int = DEFAULT_MAX_ROWS) -> ParsedFile:
    text = _decode_csv(content)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=_detect_delimiter(text))
    mapping: Optional[dict[str, int]] = None
    rows: list[RawRow] = []
    previous_end = 0
    try:
        for cells in reader:
            start = previous_end + 1
            previous_end = reader.line_num
            if not any(c.strip() for c in cells):
                continue
            if mapping is None:
                mapping = _map_header(cells)
                continue
            if len(rows) >= max_rows:
                raise ParseError(ERRO_LIMITE_LINHAS, "Arquivo acima do limite de linhas")
            rows.append(_build_row(start, cells, mapping))
    except csv.Error as ex:
        raise ParseError(ERRO_ILEGIVEL, "Arquivo CSV ilegível") from ex
    if mapping is None or not rows:
        raise ParseError(ERRO_VAZIO, "Arquivo sem linhas de dados")
    return ParsedFile(rows=tuple(rows))


def _check_zip(content: bytes, max_uncompressed: int) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > XLSX_MAX_ENTRIES:
                raise ParseError(ERRO_ILEGIVEL, "Arquivo XLSX inválido")
            if sum(i.file_size for i in infos) > max_uncompressed:
                raise ParseError(ERRO_LIMITE_BYTES, "Arquivo XLSX descomprimido grande demais")
    except zipfile.BadZipFile as ex:
        raise ParseError(ERRO_ILEGIVEL, "Arquivo XLSX ilegível") from ex


def _xlsx_text(value: Any) -> tuple[str, bool]:
    """Devolve (texto, era_numero). Número inteiro não ganha '.0'."""
    if isinstance(value, bool):
        return str(value), False
    if isinstance(value, int):
        return str(value), True
    if isinstance(value, float):
        return (str(int(value)) if value.is_integer() else repr(value)), True
    return ("" if value is None else str(value)), False


def parse_xlsx(
    content: bytes,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_uncompressed: int = DEFAULT_XLSX_MAX_UNCOMPRESSED,
) -> ParsedFile:
    _check_zip(content, max_uncompressed)
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
    except ParseError:
        raise
    except Exception as ex:
        raise ParseError(ERRO_ILEGIVEL, "Arquivo XLSX ilegível") from ex

    warnings: list[str] = []
    rows: list[RawRow] = []
    mapping: Optional[dict[str, int]] = None
    try:
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            raise ParseError(ERRO_VAZIO, "Arquivo sem linhas de dados")
        for line, cells in enumerate(
            sheet.iter_rows(max_col=XLSX_MAX_COLUMNS, values_only=True), start=1
        ):
            if line > XLSX_MAX_SCAN_LINES:
                raise ParseError(ERRO_LIMITE_LINHAS, "Arquivo acima do limite de linhas")
            cells = list(cells)
            if not any(c is not None and str(c).strip() for c in cells):
                continue
            if mapping is None:
                mapping = _map_header(cells)
                continue
            if len(rows) >= max_rows:
                raise ParseError(ERRO_LIMITE_LINHAS, "Arquivo acima do limite de linhas")
            username = None
            index = mapping["usuario"]
            if index < len(cells):
                username, numeric = _xlsx_text(cells[index])
                if numeric and WARNING_USERNAME_NUMERICO_XLSX not in warnings:
                    warnings.append(WARNING_USERNAME_NUMERICO_XLSX)
            texts = [_xlsx_text(c)[0] for c in cells]
            rows.append(_build_row(line, texts, mapping, username))
    except ParseError:
        raise
    except Exception as ex:
        raise ParseError(ERRO_ILEGIVEL, "Arquivo XLSX ilegível") from ex
    finally:
        workbook.close()
    if mapping is None or not rows:
        raise ParseError(ERRO_VAZIO, "Arquivo sem linhas de dados")
    return ParsedFile(rows=tuple(rows), warnings=tuple(warnings))


def parse_upload(
    filename: str,
    content: bytes,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_bytes: int = DEFAULT_MAX_BYTES,
) -> ParsedFile:
    name = (filename or "").lower()
    if not name.endswith((".csv", ".xlsx")):
        raise ParseError(ERRO_EXTENSAO, "Extensão inválida; use .csv ou .xlsx")
    if len(content) > max_bytes:
        raise ParseError(ERRO_LIMITE_BYTES, "Arquivo acima do limite de tamanho")
    if not content.strip():
        raise ParseError(ERRO_VAZIO, "Arquivo vazio")
    if name.endswith(".xlsx"):
        return parse_xlsx(content, max_rows=max_rows)
    return parse_csv(content, max_rows=max_rows)
