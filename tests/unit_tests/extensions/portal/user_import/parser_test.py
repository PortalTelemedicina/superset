# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.
import io
import zipfile

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from superset.extensions.portal.services.user_import import parser
from superset.extensions.portal.services.user_import.models import (
    MAX_CELL_LEN,
    WARNING_USERNAME_NUMERICO_XLSX,
)
from superset.extensions.portal.services.user_import.parser import (
    ParseError,
    parse_csv,
    parse_upload,
    parse_xlsx,
)

CSV = "NOME,USUÁRIO,SENHA,E-MAIL\nMARIA SILVA,01234567890,Senha1234,m@ptm.com.br\n"
CABECALHO_XLSX = ["NOME", "USUARIO", "SENHA", "EMAIL"]


def _xlsx(rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_csv_basico_preserva_zeros_como_texto() -> None:
    parsed = parse_csv(CSV.encode("utf-8"))
    assert len(parsed.rows) == 1
    row = parsed.rows[0]
    assert row.usuario == "01234567890"
    assert row.nome == "MARIA SILVA"
    assert row.email == "m@ptm.com.br"
    assert row.line == 2
    assert parsed.warnings == ()


def test_csv_com_bom() -> None:
    parsed = parse_csv(b"\xef\xbb\xbf" + CSV.encode("utf-8"))
    assert parsed.rows[0].nome == "MARIA SILVA"


def test_csv_ponto_e_virgula() -> None:
    parsed = parse_csv(CSV.replace(",", ";").encode("utf-8"))
    assert parsed.rows[0].usuario == "01234567890"
    assert parsed.rows[0].email == "m@ptm.com.br"


def test_csv_cp1252_fallback() -> None:
    texto = "NOME,USUÁRIO,SENHA,EMAIL\nJOÃO,1,Sênha123,j@ptm.com.br\n"
    parsed = parse_csv(texto.encode("cp1252"))
    assert parsed.rows[0].nome == "JOÃO"
    assert parsed.rows[0].senha.reveal() == "Sênha123"


def test_csv_encoding_invalido() -> None:
    with pytest.raises(ParseError) as exc:
        parse_csv(b"NOME,USUARIO,SENHA,EMAIL\n\x81\x8d,1,2,3\n")
    assert exc.value.code == parser.ERRO_ENCODING


@pytest.mark.parametrize(
    "header",
    [
        "nome,usuario,senha,email",
        "Nome , Usuário , Senha , E-mail",
        "NOME,CPF,SENHA,EMAIL",
        "NOME,USERNAME,SENHA,E MAIL",
    ],
)
def test_csv_cabecalhos_e_sinonimos(header: str) -> None:
    parsed = parse_csv(
        (header + "\nAna Lima,123,Senha1234,a@ptm.com.br\n").encode("utf-8")
    )
    row = parsed.rows[0]
    assert (row.nome, row.usuario, row.senha.reveal(), row.email) == (
        "Ana Lima",
        "123",
        "Senha1234",
        "a@ptm.com.br",
    )


def test_csv_colunas_em_outra_ordem() -> None:
    parsed = parse_csv(b"e-mail,senha,usu\xc3\xa1rio,nome\na@p.com,Senha1234,7,Ana Lima\n")
    row = parsed.rows[0]
    assert (row.nome, row.usuario, row.email) == ("Ana Lima", "7", "a@p.com")


def test_csv_coluna_role_opcional_e_perfil() -> None:
    parsed = parse_csv(
        b"NOME,USUARIO,SENHA,EMAIL,PERFIL\nA B,1,Senha1234,a@p.com,Alpha\n"
    )
    assert parsed.rows[0].role == "Alpha"
    assert parse_csv(CSV.encode()).rows[0].role == ""


def test_csv_coluna_desconhecida_ignorada() -> None:
    parsed = parse_csv(b"NOME,X,USUARIO,SENHA,EMAIL\nA B,zzz,1,Senha1234,a@p.com\n")
    assert parsed.rows[0].usuario == "1"


def test_csv_coluna_obrigatoria_ausente() -> None:
    with pytest.raises(ParseError) as exc:
        parse_csv(b"NOME,USUARIO,SENHA\nA B,1,x\n")
    assert exc.value.code == parser.ERRO_COLUNAS
    assert exc.value.missing_columns == ("EMAIL",)


def test_csv_linhas_vazias_ignoradas_e_numeracao_fisica() -> None:
    texto = (
        "NOME,USUARIO,SENHA,EMAIL\n"
        "\n"
        "A B,1,Senha1234,a@p.com\n"
        ",,,\n"
        "C D,2,Senha1234,c@p.com\n"
    )
    parsed = parse_csv(texto.encode())
    assert [r.line for r in parsed.rows] == [3, 5]


def test_csv_campo_com_quebra_de_linha_conta_linhas_fisicas() -> None:
    texto = (
        'NOME,USUARIO,SENHA,EMAIL\n"A\nB",1,Senha1234,a@p.com\n'
        "C D,2,Senha1234,c@p.com\n"
    )
    parsed = parse_csv(texto.encode())
    assert [r.line for r in parsed.rows] == [2, 4]


def test_csv_limite_de_linhas() -> None:
    corpo = "".join(f"A B,{i},Senha1234,a{i}@p.com\n" for i in range(4))
    dados = ("NOME,USUARIO,SENHA,EMAIL\n" + corpo).encode()
    assert len(parse_csv(dados, max_rows=4).rows) == 4
    with pytest.raises(ParseError) as exc:
        parse_csv(dados, max_rows=3)
    assert exc.value.code == parser.ERRO_LIMITE_LINHAS


def test_csv_so_cabecalho_e_vazio() -> None:
    with pytest.raises(ParseError) as exc:
        parse_csv(b"NOME,USUARIO,SENHA,EMAIL\n")
    assert exc.value.code == parser.ERRO_VAZIO


def test_csv_celula_longa_marca_a_linha() -> None:
    longa = "x" * (MAX_CELL_LEN + 1)
    parsed = parse_csv(
        f"NOME,USUARIO,SENHA,EMAIL\n{longa},1,Senha1234,a@p.com\n".encode()
    )
    assert parsed.rows[0].celula_longa is True
    assert len(parsed.rows[0].nome) == MAX_CELL_LEN


def test_csv_linha_curta_vira_campos_vazios() -> None:
    parsed = parse_csv(b"NOME,USUARIO,SENHA,EMAIL\nA B,1\n")
    assert parsed.rows[0].email == ""
    assert parsed.rows[0].senha.reveal() == ""


def test_xlsx_basico_com_texto_nao_gera_aviso() -> None:
    dados = _xlsx([CABECALHO_XLSX, ["A B", "01234567890", "Senha1234", "a@p.com"]])
    parsed = parse_xlsx(dados)
    assert parsed.rows[0].usuario == "01234567890"
    assert parsed.rows[0].line == 2
    assert parsed.warnings == ()


def test_xlsx_username_numerico_gera_aviso() -> None:
    dados = _xlsx(
        [["NOME", "CPF", "SENHA", "EMAIL"], ["A B", 1234567890, "Senha1234", "a@p.com"]]
    )
    parsed = parse_xlsx(dados)
    assert parsed.rows[0].usuario == "1234567890"
    assert parsed.warnings == (WARNING_USERNAME_NUMERICO_XLSX,)


def test_xlsx_float_inteiro_sem_ponto_zero() -> None:
    dados = _xlsx(
        [["NOME", "CPF", "SENHA", "EMAIL"], ["A B", 12.0, "Senha1234", "a@p.com"]]
    )
    assert parse_xlsx(dados).rows[0].usuario == "12"


def test_xlsx_linhas_vazias_e_numeracao() -> None:
    dados = _xlsx([CABECALHO_XLSX, [], ["A B", "1", "Senha1234", "a@p.com"]])
    assert [r.line for r in parse_xlsx(dados).rows] == [3]


def test_xlsx_cabecalho_ausente() -> None:
    with pytest.raises(ParseError) as exc:
        parse_xlsx(_xlsx([["X", "Y"], ["1", "2"]]))
    assert exc.value.code == parser.ERRO_COLUNAS
    assert set(exc.value.missing_columns) == {"NOME", "USUARIO", "SENHA", "EMAIL"}


def test_xlsx_limite_de_linhas() -> None:
    linhas = [CABECALHO_XLSX] + [
        ["A B", str(i), "Senha1234", f"a{i}@p.com"] for i in range(5)
    ]
    with pytest.raises(ParseError) as exc:
        parse_xlsx(_xlsx(linhas), max_rows=4)
    assert exc.value.code == parser.ERRO_LIMITE_LINHAS


def test_xlsx_zip_bomb_por_tamanho_descomprimido() -> None:
    base = _xlsx([CABECALHO_XLSX, ["A B", "1", "Senha1234", "a@p.com"]])
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as src, zipfile.ZipFile(
        buf, "w", zipfile.ZIP_DEFLATED
    ) as dst:
        for info in src.infolist():
            dst.writestr(info.filename, src.read(info.filename))
        dst.writestr("xl/bomba.bin", b"\0" * (21 * 1024 * 1024))
    bomba = buf.getvalue()
    assert len(bomba) < 100 * 1024
    with pytest.raises(ParseError) as exc:
        parse_xlsx(bomba)
    assert exc.value.code == parser.ERRO_LIMITE_BYTES


def test_xlsx_zip_com_entradas_demais() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for i in range(parser.XLSX_MAX_ENTRIES + 1):
            z.writestr(f"f{i}.txt", b"x")
    with pytest.raises(ParseError) as exc:
        parse_xlsx(buf.getvalue())
    assert exc.value.code == parser.ERRO_ILEGIVEL


def test_xlsx_ilegivel() -> None:
    with pytest.raises(ParseError) as exc:
        parse_xlsx(b"isso nao e um zip")
    assert exc.value.code == parser.ERRO_ILEGIVEL


def test_xlsx_zip_valido_sem_planilha() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("a.txt", b"x")
    with pytest.raises(ParseError) as exc:
        parse_xlsx(buf.getvalue())
    assert exc.value.code == parser.ERRO_ILEGIVEL


def test_upload_extensao_invalida_e_limites() -> None:
    with pytest.raises(ParseError) as exc:
        parse_upload("x.txt", b"abc")
    assert exc.value.code == parser.ERRO_EXTENSAO
    with pytest.raises(ParseError) as exc:
        parse_upload("x.csv", CSV.encode(), max_bytes=10)
    assert exc.value.code == parser.ERRO_LIMITE_BYTES
    with pytest.raises(ParseError) as exc:
        parse_upload("x.csv", b"  \n")
    assert exc.value.code == parser.ERRO_VAZIO


def test_upload_despacha_por_extensao_sem_diferenciar_caixa() -> None:
    assert parse_upload("USUARIOS.CSV", CSV.encode()).rows[0].line == 2
    dados = _xlsx([CABECALHO_XLSX, ["A B", "1", "Senha1234", "a@p.com"]])
    assert parse_upload("u.xlsx", dados).rows[0].usuario == "1"


def test_erro_de_parse_nao_ecoa_conteudo() -> None:
    segredo = "SenhaSecreta99"
    with pytest.raises(ParseError) as exc:
        parse_csv(f"NOME,USUARIO,SENHA\nA B,1,{segredo}\n".encode())
    assert segredo not in exc.value.message
    assert segredo not in repr(exc.value)


def test_xlsx_celula_formatada_abaixo_dos_dados_nao_estoura_limite() -> None:
    """Formatação residual longe dos dados não pode virar `limite_linhas`."""
    wb = Workbook()
    ws = wb.active
    ws.append(CABECALHO_XLSX)
    ws.append(["A B", "01234567890", "Senha1234", "a@p.com"])
    ws.cell(row=20000, column=1).font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_xlsx(buf.getvalue(), max_rows=500)
    assert [r.line for r in parsed.rows] == [2]


def test_xlsx_limite_de_linhas_conta_so_linhas_de_dados() -> None:
    linhas = [CABECALHO_XLSX]
    for i in range(3):
        linhas.append(["A B", str(i), "Senha1234", f"a{i}@p.com"])
        linhas.extend([[]] * 50)
    parsed = parse_xlsx(_xlsx(linhas), max_rows=3)
    assert len(parsed.rows) == 3
